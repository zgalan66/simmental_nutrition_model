from openpyxl import load_workbook

# data_only=False 会返回公式本身
wb = load_workbook('data/原料含量.xlsx', data_only=False)
ws = wb.active

print("=== data_only=False（看公式）===")
for r in range(3, 8):
    name = ws.cell(r, 1).value
    price = ws.cell(r, 54).value
    print(f"  行{r}: {name} => {price!r}")