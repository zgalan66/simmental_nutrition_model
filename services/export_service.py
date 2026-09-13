"""配方导出为 Excel"""
from io import BytesIO
from typing import Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from models.formula_models import Formula
from models.feed_models import Feed

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
BORDER = Border(*[Side(style="thin", color="CCCCCC")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")

NUTRIENT_LABELS = {
    "cp_pct": ("粗蛋白 CP", "%DM"),
    "ndf_pct": ("中性洗涤纤维 NDF", "%DM"),
    "adf_pct": ("酸性洗涤纤维 ADF", "%DM"),
    "me_mj_kg": ("代谢能 ME", "MJ/kgDM"),
    "tdn_pct": ("总可消化养分 TDN", "%DM"),
    "ca_pct": ("钙 Ca", "%DM"),
    "p_pct": ("磷 P", "%DM"),
    "dm_pct": ("干物质 DM", "%"),
}


async def export_formula_excel(db: AsyncSession, formula_id: str) -> bytes:
    f = await db.get(Formula, formula_id)
    if not f:
        return None

    wb = Workbook()

    # ===== Sheet 1: 配方详情 =====
    ws = wb.active
    ws.title = "配方详情"

    ws["A1"] = f"配方：{f.name or '未命名'}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 24

    ws["A2"] = "配方 ID"
    ws["B2"] = f.id
    ws["A3"] = "创建时间"
    ws["B3"] = f.created_at.strftime("%Y-%m-%d %H:%M") if f.created_at else ""
    ws["A4"] = "总成本（元/单位鲜重）"
    ws["B4"] = f.total_cost
    ws["A5"] = "每 kg 干物质成本"
    cost_dm = None
    if f.total_cost and f.nutrition_result and f.nutrition_result.get("dm_pct"):
        dm = f.nutrition_result["dm_pct"] / 100
        if dm > 0:
            cost_dm = round(f.total_cost / dm, 4)
    ws["B5"] = cost_dm

    for row in ws.iter_rows(min_row=2, max_row=5, min_col=1, max_col=1):
        for c in row:
            c.font = Font(bold=True)

    # ===== Sheet 2: 原料明细 =====
    ws2 = wb.create_sheet("原料明细")

    headers = ["原料名", "比例（%）", "单价（元/kg）", "DM%", "CP%", "NDF%", "ADF%", "ME(MJ/kg)", "TDN%", "Ca%", "P%"]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(1, col, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    ing = f.ingredients or []
    total_weight = sum(it.get("ratio", 0) for it in ing) or 1

    feed_ids = [it.get("feed_id") for it in ing if it.get("feed_id")]
    feeds_map = {}
    if feed_ids:
        r = await db.execute(select(Feed).where(Feed.id.in_(feed_ids)))
        for fd in r.scalars().all():
            feeds_map[fd.id] = fd

    row = 2
    for it in ing:
        name = it.get("name") or ""
        ratio = it.get("ratio", 0)
        fd = feeds_map.get(it.get("feed_id"))

        ws2.cell(row, 1, name).border = BORDER
        ws2.cell(row, 2, round(ratio / total_weight * 100, 2)).border = BORDER
        if fd:
            ws2.cell(row, 3, fd.price_yuan_kg).border = BORDER
            ws2.cell(row, 4, fd.dm_pct).border = BORDER
            ws2.cell(row, 5, fd.cp_pct).border = BORDER
            ws2.cell(row, 6, fd.ndf_pct).border = BORDER
            ws2.cell(row, 7, fd.adf_pct).border = BORDER
            ws2.cell(row, 8, fd.me_mj_kg).border = BORDER
            ws2.cell(row, 9, fd.tdn_pct).border = BORDER
            ws2.cell(row, 10, fd.ca_pct).border = BORDER
            ws2.cell(row, 11, fd.p_pct).border = BORDER
        row += 1

    ws2.cell(row, 1, "合计").font = Font(bold=True)
    ws2.cell(row, 2, round(total_weight / total_weight * 100, 2)).font = Font(bold=True)

    # 列宽
    widths = [22, 10, 13, 9, 9, 9, 9, 13, 9, 9, 9]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[chr(64 + i)].width = w

    # ===== Sheet 3: 营养汇总 =====
    ws3 = wb.create_sheet("营养汇总")
    ws3["A1"] = "营养指标"
    ws3["B1"] = "数值"
    ws3["C1"] = "单位"
    for c in ("A1", "B1", "C1"):
        ws3[c].fill = HEADER_FILL
        ws3[c].font = HEADER_FONT
        ws3[c].alignment = CENTER
        ws3[c].border = BORDER

    nr = f.nutrition_result or {}
    r3 = 2
    for field, (label, unit) in NUTRIENT_LABELS.items():
        v = nr.get(field)
        if v is None:
            continue
        ws3.cell(r3, 1, label).border = BORDER
        ws3.cell(r3, 2, v).border = BORDER
        ws3.cell(r3, 3, unit).border = BORDER
        r3 += 1

    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 12

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
