from openpyxl import load_workbook

wb = load_workbook('data/原料含量.xlsx', data_only=True)
ws = wb.active

print(f"表格大小: {ws.max_row} 行 x {ws.max_column} 列")
print()

# 先找出哪一列是价格：打印第 50~60 列，第 3 行的值
print("第 3 行的第 50~60 列：")
for c in range(50, 61):
    v = ws.cell(3, c).value
    print(f"  第 {c} 列 ({ws.cell(2, c).coordinate}): {v!r}")

print()
print("前 10 行 A 列(名称) 和 BB 列(第 54 列)：")
for r in range(3, 13):
    name = ws.cell(r, 1).value
    price = ws.cell(r, 54).value
    print(f"  行{r}: {name} => {price!r}")